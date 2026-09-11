import os
import pandas as pd
from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def load_document(file_path):

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return read_pdf(file_path)

    elif extension == ".docx":
        return read_docx(file_path)

    elif extension == ".txt":
        return read_txt(file_path)

    elif extension == ".csv":
        return read_csv(file_path)

    elif extension == ".xlsx":
        return read_xlsx(file_path)

    elif extension == ".pptx":
        return read_pptx(file_path)

    else:
        raise ValueError(f"Unsupported file type: {extension}")


def read_pdf(path):

    reader = PdfReader(path)

    text = ""

    for page in reader.pages:
        text += page.extract_text() or ""
        text += "\n"

    return text


def read_docx(path):

    doc = Document(path)

    return "\n".join(
        paragraph.text
        for paragraph in doc.paragraphs
    )


def read_txt(path):

    with open(path, "r", encoding="utf-8") as file:
        return file.read()


def read_csv(path):

    df = pd.read_csv(path)

    return df.to_string(index=False)


def read_xlsx(path):

    df = pd.read_excel(path)

    return df.to_string(index=False)


def read_pptx(path):

    prs = Presentation(path)

    text = ""

    for slide in prs.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text += shape.text + "\n"

    return text